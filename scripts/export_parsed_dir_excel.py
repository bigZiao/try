import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook


SUMMARY_HEADERS = [
    "file",
    "status",
    "parse_status",
    "need_review",
    "review_reason",
    "bill_type",
    "merchant_name",
    "order_date",
    "customer_name",
    "customer_phone",
    "order_no",
    "salesperson",
    "item_count",
    "total_quantity",
    "total_amount",
    "warnings",
    "validation_errors",
]

ITEM_HEADERS = [
    "file",
    "row_no",
    "style_no",
    "product_name",
    "color",
    "size",
    "quantity",
    "unit_price",
    "subtotal",
    "block_indexes",
]

BILL_ITEM_HEADERS = [
    "source_file",
    "status",
    "need_review",
    "merchant_name",
    "order_date",
    "customer_name",
    "style_no",
    "product_name",
    "color",
    "size",
    "quantity",
    "unit_price",
    "subtotal",
    "bill_total_quantity",
    "bill_total_amount",
    "payments",
    "warnings",
    "confidence",
    "block_indexes",
]


def main() -> None:
    parser = argparse.ArgumentParser(description="Export parsed OCR JSON results to Excel.")
    parser.add_argument("input_dir", type=Path)
    parser.add_argument("--output", type=Path, default=Path("data/parsed_ocr_summary.xlsx"))
    args = parser.parse_args()

    files = sorted(args.input_dir.glob("*.parsed.json"))
    if not files:
        raise SystemExit(f"No parsed JSON files found in {args.input_dir}")

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "summary"
    bill_items_sheet = workbook.create_sheet("bill_items")
    items_sheet = workbook.create_sheet("items")

    summary_sheet.append(SUMMARY_HEADERS)
    bill_items_sheet.append(BILL_ITEM_HEADERS)
    items_sheet.append(ITEM_HEADERS)

    for file_path in files:
        payload = json.loads(file_path.read_text(encoding="utf-8"))
        data = payload.get("final_json") or payload.get("corrected_json") or payload.get("structured_json") or {}
        summary = data.get("summary") if isinstance(data.get("summary"), dict) else {}
        items = data.get("items") if isinstance(data.get("items"), list) else []

        append_row(
            summary_sheet,
            [
                file_path.name,
                payload.get("status"),
                get_value(data, "parse_status"),
                get_value(data, "need_review"),
                get_value(data, "review_reason"),
                get_value(data, "bill_type"),
                get_value(data, "merchant_name", "商家", "店铺"),
                normalize_date(get_value(data, "order_date", "日期")),
                get_value(data, "customer_name", "客户", "客户名称"),
                get_value(data, "customer_phone", "电话", "客户电话"),
                get_value(data, "order_no", "订单号", "批次"),
                get_value(data, "salesperson", "店员", "销售员"),
                len(items),
                get_value(summary, "total_quantity", "quantity_total", "总数量"),
                get_value(summary, "total_amount", "总金额"),
                compact(data.get("warnings")),
                compact(payload.get("validation_errors")),
            ]
        )

        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                continue
            for size_name, quantity in expand_size_rows(item):
                append_row(
                    bill_items_sheet,
                    [
                        file_path.name,
                        payload.get("status"),
                        get_value(data, "need_review"),
                        get_value(data, "merchant_name", "商家", "店铺"),
                        normalize_date(get_value(data, "order_date", "日期")),
                        get_value(data, "customer_name", "客户", "客户名称"),
                        get_value(item, "style_no", "style_code", "sku_id", "sku_code", "item_code", "item_no", "product_code", "style_number", "spu_id", "goods_code", "款号"),
                        get_value(item, "product_name", "name", "item_name", "product_title", "goods_name", "名称", "商品名称"),
                        get_value(item, "color", "color_name", "颜色"),
                        size_name,
                        quantity,
                        get_value(item, "unit_price", "单价"),
                        calc_subtotal_for_quantity(item, quantity),
                        get_value(summary, "total_quantity", "quantity_total", "总数量"),
                        get_value(summary, "total_amount", "总金额"),
                        compact(data.get("payments")),
                        compact(data.get("warnings")),
                        get_value(data, "confidence"),
                        compact(get_value(item, "block_indexes")),
                    ]
                )
            append_row(
                items_sheet,
                [
                    file_path.name,
                    index,
                    get_value(item, "style_no", "style_code", "sku_id", "sku_code", "item_code", "item_no", "product_code", "style_number", "spu_id", "goods_code", "款号"),
                    get_value(item, "product_name", "name", "item_name", "product_title", "goods_name", "名称", "商品名称"),
                    get_value(item, "color", "color_name", "颜色"),
                    normalize_size(get_value(item, "size", "size_name", "尺码")),
                    get_value(item, "quantity", "total_quantity", "qty", "数量"),
                    get_value(item, "unit_price", "单价"),
                    get_value(item, "subtotal", "amount", "小计", "金额"),
                    compact(get_value(item, "block_indexes")),
                ]
            )

    autosize(summary_sheet)
    autosize(bill_items_sheet)
    autosize(items_sheet)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(args.output)


def get_value(data: dict[str, Any], *keys: str, default: Any = None) -> Any:
    for key in keys:
        if key in data and data[key] is not None:
            return data[key]
    return default


def compact(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def append_row(sheet, row: list[Any]) -> None:
    sheet.append([excel_value(value) for value in row])


def excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return compact(value)
    return value


def normalize_date(value: Any) -> Any:
    if value is None:
        return None

    text = str(value).strip()
    match = re.search(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        return text

    year, month, day = match.groups()
    return f"{year}-{int(month):02d}-{int(day):02d}"


def normalize_size(value: Any) -> Any:
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d+(?:\.\d+)?", text):
        return None
    if ":" in text or "：" in text or "," in text or "，" in text:
        return None

    valid_sizes = {"均码", "小", "中", "大", "S", "M", "L", "XL", "XXL", "XXXL", "XXXXL", "XS", "XXS", "F", "FREE"}
    if text in valid_sizes or text.upper() in valid_sizes:
        return text
    return text


def expand_size_rows(item: dict[str, Any]) -> list[tuple[Any, Any]]:
    sizes = get_value(item, "sizes", "size_quantities", "尺码明细")
    if isinstance(sizes, dict):
        rows = []
        for size_name, quantity in sizes.items():
            normalized = normalize_size(size_name)
            if normalized and to_number(quantity):
                rows.append((normalized, quantity))
        if rows:
            return rows

    if isinstance(sizes, list):
        rows = []
        for size_item in sizes:
            if not isinstance(size_item, dict):
                continue
            size_name = normalize_size(get_value(size_item, "size_name", "size", "尺码"))
            quantity = get_value(size_item, "quantity", "count", "qty", "num", "数量")
            if size_name and to_number(quantity):
                rows.append((size_name, quantity))
        if rows:
            return rows

    return [
        (
            normalize_size(get_value(item, "size", "size_name", "尺码")),
            get_value(item, "quantity", "total_quantity", "qty", "数量"),
        )
    ]


def calc_subtotal_for_quantity(item: dict[str, Any], quantity: Any) -> Any:
    unit_price = to_number(get_value(item, "unit_price", "单价"))
    size_quantity = to_number(quantity)
    if unit_price is not None and size_quantity is not None:
        return unit_price * size_quantity
    return get_value(item, "subtotal", "amount", "小计", "金额")


def to_number(value: Any) -> float | int | None:
    if value in (None, ""):
        return None
    try:
        number = float(str(value).replace(",", ""))
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def autosize(sheet) -> None:
    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 60)


if __name__ == "__main__":
    main()
